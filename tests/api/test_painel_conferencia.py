"""Painel do mês e conferência (cartões, seleção de empresas, Excel)."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from .conftest import API


def _linha(tipo_id, qa, va, qc=None, vc=None):
    return {"tipoId": tipo_id, "qtdeAutor": qa, "valorAutor": va, "qtdeCanc": qc, "valorCanc": vc}


def _preparar(entrar, nova_empresa, tipos):
    a = nova_empresa("Alfa", tipos_codigos=("NFE_ENT", "NFE_SAI", "NFCE"))
    b = nova_empresa("Beta", tipos_codigos=("NFCE",))
    h = nova_empresa("Da Hellen", responsavel="Hellen", tipos_codigos=("NFCE",))
    c = entrar("Natan")
    c.put(f"{API}/lancamento/{a}/2026-08", json={"linhas": [_linha(tipos["NFE_ENT"], 2, "100"),
                                                             _linha(tipos["NFE_SAI"], 3, "300,50", 1, 10)]})
    c.put(f"{API}/fechamento/{b}/2026-08", json={"status": "sem_movimento"})
    return c, a, b, h


def test_painel_contagem_progresso_e_filtro(entrar, nova_empresa, tipos):
    c, a, b, _h = _preparar(entrar, nova_empresa, tipos)
    j = c.get(f"{API}/painel?competencia=2026-08").json()
    natan = c.get(f"{API}/sessao").json()["usuario"]["id"]
    assert j["total"] == 2 and j["feitos"] == 2 and j["progresso"] == 100 and j["carteira"] == str(natan)
    assert j["contagem"] == {"pendente": 0, "preenchido": 1, "semMovimento": 1}
    alfa = next(e for e in j["empresas"] if e["empresa"]["id"] == a)
    assert alfa["status"] == "preenchido" and alfa["entradas"] == "100.00" and alfa["saidas"] == "300.50"
    assert alfa["tiposUsados"] == 3 and alfa["tiposLancados"] == 2
    j = c.get(f"{API}/painel?competencia=2026-08&status=sem_movimento").json()
    assert [e["empresa"]["nome"] for e in j["empresas"]] == ["Beta"] and j["filtroStatus"] == "sem_movimento"
    j = c.get(f"{API}/painel?competencia=2026-08&carteira=todos").json()
    assert j["total"] == 3 and j["somenteLeitura"] is True


def test_conferencia_cartoes_blocos_e_bateu(entrar, nova_empresa, tipos):
    c, a, _b, _h = _preparar(entrar, nova_empresa, tipos)
    j = c.get(f"{API}/conferencia?competencia=2026-08").json()
    assert j["progresso"] == {"linhas": 2, "bateram": 0, "percentual": 0, "empresasOk": 0, "empresasPendentes": 1}
    assert j["somenteLeitura"] is False and j["nomeCarteira"] == "Natan" and j["haLancamentos"] is True
    (cartao,) = j["cartoes"]
    assert cartao["empresa"]["nome"] == "Alfa" and cartao["podeEditar"] is True and cartao["faltam"] == 2
    assert [b["titulo"] for b in cartao["blocos"]] == ["ENTRADAS", "SAÍDAS"]
    sai = cartao["blocos"][1]
    assert sai["subtotal"] == {"qtdeAutor": 3, "valorAutor": "300.50", "qtdeCanc": 1, "valorCanc": "10.00"}
    lid = sai["linhas"][0]["id"]
    c.put(f"{API}/lancamentos/{lid}/bateu", json={"bateu": True})
    j = c.get(f"{API}/conferencia?competencia=2026-08&vista=pendencia").json()
    assert j["progresso"]["bateram"] == 1 and j["vista"] == "pendencia" and len(j["cartoes"]) == 1
    c.put(f"{API}/conferencia/{a}/2026-08/bateu", json={"bateu": True})
    j = c.get(f"{API}/conferencia?competencia=2026-08&vista=pendencia").json()
    assert j["cartoes"] == [] and j["haLancamentos"] is True and j["progresso"]["empresasOk"] == 1


def test_conferencia_selecao_de_empresas_e_outra_carteira(entrar, nova_empresa, tipos):
    c, a, _b, h = _preparar(entrar, nova_empresa, tipos)
    c2 = entrar("Hellen")
    c2.put(f"{API}/lancamento/{h}/2026-08", json={"linhas": [_linha(tipos["NFCE"], 9, "90")]})
    c = entrar("Natan")
    j = c.get(f"{API}/conferencia?competencia=2026-08&carteira=todos").json()
    assert {o["nome"] for o in j["opcoes"]} == {"Alfa", "Da Hellen"} and j["somenteLeitura"] is True
    assert all(not ct["podeEditar"] for ct in j["cartoes"])
    j = c.get(f"{API}/conferencia?competencia=2026-08&carteira=todos&empresa={h}").json()
    assert [ct["empresa"]["nome"] for ct in j["cartoes"]] == ["Da Hellen"] and j["escolhidas"] == [h]
    assert j["cartoes"][0]["responsavel"] == "Hellen"
    # seleção de outra carteira é ignorada (mostra todas da carteira escolhida)
    j = c.get(f"{API}/conferencia?competencia=2026-08&empresa={h}").json()
    assert [ct["empresa"]["nome"] for ct in j["cartoes"]] == ["Alfa"] and j["escolhidas"] == []


def test_conferencia_excel(entrar, nova_empresa, tipos):
    c, a, _b, _h = _preparar(entrar, nova_empresa, tipos)
    r = c.get(f"{API}/conferencia/excel?competencia=2026-08")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="conferencia-2026-08.xlsx"'
    ws = load_workbook(io.BytesIO(r.content)).active
    linhas = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert linhas[0][:6] == ["Alfa", "Natan", "Entrada", "NF-e Entradas", 2, 100]


def test_referencias(entrar):
    j = entrar("Natan").get(f"{API}/referencias?competencia=2026-08").json()
    assert j["competencia"]["extenso"] == "Agosto de 2026" and len(j["competencias"]) >= 18
    assert [u["nome"] for u in j["usuarios"]] == ["Hellen", "Jaísa", "Mariana", "Natan"]
    assert len(j["tipos"]) == 7 and "Importado da planilha" not in j["motivosAjuste"]
    assert {"valor": "sem_movimento", "rotulo": "Sem movimento"} in j["status"]

"""Importação da planilha antiga (FATURAMENTO_V2). Planilhas sintéticas — nenhum dado de cliente."""
from __future__ import annotations

import io
import json
from datetime import datetime
from decimal import Decimal

import pytest

from conftest import TOKEN_CSRF
from openpyxl import Workbook
from sqlalchemy import select

from faturamento.models import (
    ORIGEM_PLANILHA, Ajuste, Auditoria, Empresa, Importacao, Lancamento,
)
from faturamento.services import RegraViolada
from faturamento.services import importacao as svc
from faturamento.util import ValorInvalido, competencia

CABECALHO = ["EMPRESA", "COMP.", "TIPO", "QTDE_AUTOR", "VALOR_AUTOR", "AJUSTE_QTDE_AUTOR", "AJUSTE_VALOR_AUTOR",
             "TOTAL_QTDE_AUTOR", "TOTAL_VALOR_AUTOR", "QTDE_CANC", "VALOR_CANC", "AJUSTE_QTDE_CANC",
             "AJUSTE_VALOR_CANC", "TOTAL_QTDE_CANC", "TOTAL_VALOR_CANC", "EXIBIR_DINAMICA"]
JAN = datetime(2026, 1, 1)
FEV = datetime(2026, 2, 1)


def _linha(emp, comp, tipo, qa=None, va=None, aqa=None, ava=None, qc=None, vc=None, aqc=None, avc=None):
    return [emp, comp, tipo, qa, va, aqa, ava, None, None, qc, vc, aqc, avc, None, None, "SIM"]


LINHAS_PADRAO = [
    _linha("LOJA ALFA (MATRIZ)", JAN, "NFC-e", 100, 1000.5, 5, 50.25, 2, 20),            # linha 3
    _linha("LOJA ALFA (MATRIZ)", JAN, "NF-e ENTRADAS", 10, "=1000+234.56"),               # linha 4: conta digitada
    _linha("LOJA ALFA (MATRIZ)", JAN, "NF-e SAÍDAS"),                                     # linha 5: vazia
    _linha("PADARIA BETA", JAN, "NF-e ENTRADAS", None, 500),                              # linha 6: valor sem qtde
    _linha("PADARIA BETA", FEV, "NFS-e SAÍDAS", 3, 300, None, None, None, None, -1, -10),  # linha 7: ajuste canc
    _linha("PADARIA BETA", FEV, "MDF-e", 1, 10),                                          # linha 8: tipo desconhecido
    _linha("LOJA ALFA (MATRIZ)", JAN, "NFC-e", 999, 9999),                                # linha 9: repetida
    _linha("PADARIA BETA", "xx", "NFC-e", 1, 1),                                          # linha 10: data inválida
    _linha("PADARIA BETA", FEV, "CT-e ENTRADAS"),                                         # linha 11: vazia
]


def montar_planilha(linhas=None, *, linha_cab=2, colunas=None, aba="DADOS") -> bytes:
    """Reproduz o layout da FATURAMENTO_V2: título vazio na linha 1, cabeçalho em B2..Q2."""
    linhas = LINHAS_PADRAO if linhas is None else linhas
    wb = Workbook()
    wb.active.title = "VALID. DE DADOS"
    wb.active["B2"] = "VALIDAÇÃO DE DADOS 1"
    ws = wb.create_sheet(aba)
    cab = colunas or CABECALHO
    for j, nome in enumerate(cab, start=2):
        ws.cell(row=linha_cab, column=j, value=nome)
    for i, lin in enumerate(linhas, start=linha_cab + 1):
        valores = dict(zip(CABECALHO, lin))
        for j, nome in enumerate(cab, start=2):
            v = valores.get(nome)
            if nome == "TOTAL_QTDE_AUTOR":
                v = f"=E{i}+G{i}"
            elif nome == "TOTAL_VALOR_AUTOR":
                v = f"=F{i}+H{i}"
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def dados_tmp(app, tmp_path):
    app.config["DADOS_DIR"] = tmp_path
    return tmp_path


def _lanc(db, nome_emp, comp, tipo):
    e = db.session.scalar(select(Empresa).where(Empresa.nome == nome_emp))
    return db.session.scalar(select(Lancamento).where(
        Lancamento.empresa_id == e.id, Lancamento.competencia == competencia(comp), Lancamento.tipo_id == tipo.id))


# ====================================================================== analisar

def test_analisar_le_layout_e_avisos(app, db):
    p = svc.analisar(montar_planilha(), "FATURAMENTO_V2.xlsx")
    assert p.aba == "DADOS" and p.linha_cabecalho == 2
    assert p.linhas_vazias == 2
    assert p.n_lancamentos == 4
    assert p.n_ajustes == 2
    textos = [a.texto for a in p.avisos]
    assert any(t.startswith("F4: célula com conta digitada (=1000+234.56); importado o resultado") for t in textos)
    sem_qtde = next(a for a in p.avisos if "sem quantidade" in a.texto)
    assert sem_qtde.destaque and sem_qtde.linha == 6
    assert "Linha 6: valor R$ 500,00 sem quantidade" in sem_qtde.texto
    assert "NF-e Entradas, Padaria Beta, 01/2026" in sem_qtde.texto
    assert any("TIPO desconhecido 'MDF-e'" in t for t in textos)
    assert any(t.startswith("Linha 9: repete") for t in textos)
    assert any(t.startswith("Linha 10: competência inválida") for t in textos)

    itens = {(i.empresa, i.comp, i.tipo_codigo): i for i in p.itens}
    nfce = itens[("LOJA ALFA (MATRIZ)", "2026-01", "NFCE")]
    assert (nfce.qtde_autor, nfce.valor_autor, nfce.aj_qtde_autor, nfce.aj_valor_autor) == (100, "1000.50", 5, "50.25")
    assert (nfce.qtde_canc, nfce.valor_canc) == (2, "20.00")
    assert itens[("LOJA ALFA (MATRIZ)", "2026-01", "NFE_ENT")].valor_autor == "1234.56"
    beta = itens[("PADARIA BETA", "2026-01", "NFE_ENT")]
    assert (beta.qtde_autor, beta.valor_autor) == (0, "500.00")
    assert beta.qtde_canc is None and beta.valor_canc is None
    nfse = itens[("PADARIA BETA", "2026-02", "NFSE_SAI")]
    assert (nfse.aj_qtde_canc, nfse.aj_valor_canc) == (-1, "-10.00")


def test_analisar_estatisticas(app, db, nova_empresa):
    nova_empresa("Padaria Beta", tipos_codigos=("NFE_ENT",))
    p = svc.analisar(montar_planilha(), "f.xlsx")
    emps = {e["nome"]: e for e in p.empresas()}
    assert emps["Loja Alfa (Matriz)"]["nova"] is True
    assert emps["Padaria Beta"]["nova"] is False
    assert emps["Loja Alfa (Matriz)"]["n_lancamentos"] == 2
    assert [c.month for c in p.competencias] == [1, 2]
    totais = {t["competencia"].month: t for t in p.totais_por_competencia()}
    # jan: saídas = NFC-e 1000,50 + ajuste 50,25; entradas = 1234,56 + 500
    assert totais[1]["S"]["valor"] == Decimal("1050.75")
    assert totais[1]["S"]["qtde"] == 105
    assert totais[1]["E"]["valor"] == Decimal("1734.56")
    assert totais[2]["S"]["qtde_canc"] == -1


def test_cabecalho_por_nome_de_coluna(app, db):
    colunas = ["TIPO", "EMPRESA", "VALOR_AUTOR", "QTDE_AUTOR", "COMP."]
    dados = montar_planilha([_linha("LOJA ALFA", JAN, "NF-e SAÍDAS", 7, 70)], linha_cab=4, colunas=colunas, aba="Planilha1")
    p = svc.analisar(dados, "x.xlsx")
    assert p.linha_cabecalho == 4
    assert len(p.itens) == 1
    i = p.itens[0]
    assert (i.tipo_codigo, i.qtde_autor, i.valor_autor) == ("NFE_SAI", 7, "70.00")


def test_arquivo_invalido(app):
    with pytest.raises(ValorInvalido):
        svc.analisar(b"isto nao e xlsx", "x.xlsx")
    wb = Workbook()
    wb.active["A1"] = "sem cabeçalho"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValorInvalido, match="cabeçalho"):
        svc.analisar(buf.getvalue(), "x.xlsx")


def test_avaliar_conta():
    assert svc.avaliar_conta("=75536.37-11.5") == Decimal("75524.87")
    assert svc.avaliar_conta("=(10+5)*2/3") == Decimal("10")
    assert svc.avaliar_conta("=SUM(A1:A2)") is None
    assert svc.avaliar_conta("=__import__('os')") is None


def test_guardar_e_carregar_previa(app, dados_tmp):
    p = svc.analisar(montar_planilha(), "f.xlsx")
    token = svc.guardar_previa(p)
    arq = dados_tmp / "importacoes" / f"{token}.json"
    assert arq.exists()
    json.loads(arq.read_text(encoding="utf-8"))  # JSON puro, decimais como texto
    p2 = svc.carregar_previa(token)
    assert [(i.empresa, i.comp, i.tipo_codigo, i.valor_autor) for i in p2.itens] == \
           [(i.empresa, i.comp, i.tipo_codigo, i.valor_autor) for i in p.itens]
    assert [a.texto for a in p2.avisos] == [a.texto for a in p.avisos]
    with pytest.raises(ValorInvalido):
        svc.carregar_previa("../../etc/passwd")
    svc.descartar(token)
    assert not arq.exists()


# ====================================================================== confirmar

def test_confirmar_cria_empresas_lancamentos_e_ajustes(app, db, usuarios, tipos):
    p = svc.analisar(montar_planilha(), "FATURAMENTO_V2.xlsx")
    imp = svc.confirmar(p, usuarios["Mariana"], usuarios["Natan"], "FATURAMENTO_V2.xlsx", outras_carteiras=True)
    db.session.commit()

    alfa = db.session.scalar(select(Empresa).where(Empresa.nome == "Loja Alfa (Matriz)"))
    beta = db.session.scalar(select(Empresa).where(Empresa.nome == "Padaria Beta"))
    assert alfa and beta
    assert {t.codigo for t in alfa.tipos} == {"NFCE", "NFE_ENT"}
    assert {t.codigo for t in beta.tipos} == {"NFE_ENT", "NFSE_SAI"}
    vinc = alfa.historico_responsaveis[0]
    assert vinc.usuario.nome == "Mariana" and vinc.inicio == competencia("2026-01") and vinc.origem == ORIGEM_PLANILHA

    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert (l.qtde_autor, l.valor_autor, l.qtde_canc, l.valor_canc) == (100, Decimal("1000.50"), 2, Decimal("20.00"))
    assert (l.total_qtde_autor, l.total_valor_autor) == (105, Decimal("1050.75"))
    assert l.importacao_id == imp.id and l.responsavel_id == usuarios["Mariana"].id
    assert l.criado_por_id == usuarios["Natan"].id
    assert [a.motivo for a in l.ajustes] == ["Importado da planilha"]
    assert _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFE_ENT"]).valor_autor == Decimal("1234.56")
    lb = _lanc(db, "Padaria Beta", "2026-01", tipos["NFE_ENT"])
    assert (lb.qtde_autor, lb.valor_autor) == (0, Decimal("500.00"))
    ln = _lanc(db, "Padaria Beta", "2026-02", tipos["NFSE_SAI"])
    assert (ln.total_qtde_canc, ln.total_valor_canc) == (-1, Decimal("-10.00"))

    r = svc.resumo(imp)
    assert r["lancamentos_criados"] == 4 and r["ajustes_criados"] == 2
    assert sorted(r["empresas_criadas"]) == ["Loja Alfa (Matriz)", "Padaria Beta"]
    assert any("sem quantidade" in a["texto"] and a["destaque"] for a in r["avisos"])
    assert imp.responsavel_id == usuarios["Mariana"].id and imp.usuario_id == usuarios["Natan"].id
    acoes = {(a.tabela, a.acao) for a in db.session.scalars(select(Auditoria))}
    assert {("fat_importacao", "importar"), ("fat_lancamento", "criar"), ("fat_ajuste", "criar"),
            ("fat_empresa", "criar")} <= acoes


def test_reimportar_mesmo_arquivo_nao_duplica(app, db, usuarios):
    dados = montar_planilha()
    svc.confirmar(svc.analisar(dados, "f.xlsx"), usuarios["Natan"], usuarios["Natan"])
    db.session.commit()
    imp2 = svc.confirmar(svc.analisar(dados, "f.xlsx"), usuarios["Natan"], usuarios["Natan"])
    db.session.commit()
    r = svc.resumo(imp2)
    assert r["lancamentos_criados"] == 0 and r["lancamentos_alterados"] == 0 and r["ajustes_criados"] == 0
    assert r["sem_mudanca"] == 4 and r["empresas_existentes"] == 2
    assert db.session.query(Lancamento).count() == 4
    assert db.session.query(Ajuste).count() == 2
    assert db.session.query(Empresa).count() == 2


def test_existente_ganha_tipos_e_e_sobrescrito_com_auditoria(app, db, usuarios, tipos, nova_empresa):
    alfa = nova_empresa("Loja Alfa (Matriz)", responsavel="Hellen", tipos_codigos=("NFCE",))
    db.session.add(Lancamento(empresa_id=alfa.id, competencia=competencia("2026-01"), tipo_id=tipos["NFCE"].id,
                              qtde_autor=1, valor_autor=Decimal("1.00")))
    db.session.commit()
    imp = svc.confirmar(svc.analisar(montar_planilha(), "f.xlsx"), usuarios["Natan"], usuarios["Natan"],
                        outras_carteiras=True)
    db.session.commit()
    db.session.refresh(alfa)
    assert {t.codigo for t in alfa.tipos} == {"NFCE", "NFE_ENT"}
    assert alfa.historico_responsaveis[0].usuario.nome == "Hellen"  # carteira existente não muda
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert (l.qtde_autor, l.valor_autor) == (100, Decimal("1000.50"))
    assert l.importacao_id is None  # já existia: desfazer restaura, não apaga
    reg = db.session.scalar(select(Auditoria).where(Auditoria.tabela == "fat_lancamento", Auditoria.acao == "alterar",
                                                    Auditoria.registro_id == l.id))
    assert json.loads(reg.antes)["qtde_autor"] == 1 and json.loads(reg.depois)["qtde_autor"] == 100
    r = svc.resumo(imp)
    assert r["lancamentos_alterados"] == 1 and "Loja Alfa (Matriz): NF-e Entradas" in r["tipos_adicionados"]


def test_bateu_nao_e_sobrescrito(app, db, usuarios, tipos, nova_empresa):
    alfa = nova_empresa("Loja Alfa (Matriz)", tipos_codigos=("NFCE",))
    db.session.add(Lancamento(empresa_id=alfa.id, competencia=competencia("2026-01"), tipo_id=tipos["NFCE"].id,
                              qtde_autor=1, valor_autor=Decimal("1.00"), bateu=True))
    db.session.commit()
    imp = svc.confirmar(svc.analisar(montar_planilha(), "f.xlsx"), usuarios["Natan"], usuarios["Natan"])
    db.session.commit()
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert (l.qtde_autor, l.valor_autor, l.ajustes) == (1, Decimal("1.00"), [])
    r = svc.resumo(imp)
    assert r["ignorados_bateu"] == 1
    assert any("Bateu" in a["texto"] for a in r["avisos"])


def test_confirmar_sem_itens(app, usuarios):
    p = svc.analisar(montar_planilha([_linha("X", JAN, "NFC-e")]), "f.xlsx")
    with pytest.raises(RegraViolada):
        svc.confirmar(p, usuarios["Natan"], usuarios["Natan"])


# ====================================================================== desfazer

def test_desfazer(app, db, usuarios, tipos, nova_empresa):
    alfa = nova_empresa("Loja Alfa (Matriz)", tipos_codigos=("NFCE",))
    db.session.add(Lancamento(empresa_id=alfa.id, competencia=competencia("2026-01"), tipo_id=tipos["NFCE"].id,
                              qtde_autor=1, valor_autor=Decimal("1.00")))
    db.session.commit()
    imp = svc.confirmar(svc.analisar(montar_planilha(), "f.xlsx"), usuarios["Natan"], usuarios["Natan"])
    db.session.commit()
    # uma linha criada pela importação foi conferida depois: deve ficar
    conferida = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFE_ENT"])
    conferida.bateu = True
    db.session.commit()

    out = svc.desfazer(imp, usuarios["Natan"])
    db.session.commit()
    assert imp.desfeita is True
    assert out["mantidos_bateu"] == 1 and out["removidos"] == 2 and out["restaurados"] == 1
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert (l.qtde_autor, l.valor_autor, l.ajustes) == (1, Decimal("1.00"), [])
    assert db.session.scalar(select(Empresa).where(Empresa.nome == "Padaria Beta")) is None  # criada e esvaziada
    assert db.session.query(Lancamento).count() == 2
    assert db.session.scalar(select(Auditoria).where(Auditoria.tabela == "fat_importacao", Auditoria.acao == "desfazer"))
    with pytest.raises(RegraViolada):
        svc.desfazer(imp, usuarios["Natan"])


# ====================================================================== telas

def _enviar(cli, dados=None, responsavel_id=None, nome="FATURAMENTO_V2.xlsx"):
    form = {"arquivo": (io.BytesIO(dados or montar_planilha()), nome), "csrf_token": TOKEN_CSRF}
    if responsavel_id:
        form["responsavel_id"] = str(responsavel_id)
    return cli.post("/importacao/previa", data=form, content_type="multipart/form-data")


def _token(html: str) -> str:
    import re
    return re.search(r'name="token" value="([^"]+)"', html).group(1)


def test_tela_fluxo_completo(app, db, login, usuarios):
    cli = login("Natan")
    assert cli.get("/importacao").status_code == 200
    resp = _enviar(cli, responsavel_id=usuarios["Jaísa"].id)
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert "Será criada" in html and "Loja Alfa (Matriz)" in html
    assert "sem quantidade" in html and "conta digitada" in html
    assert "Totais por competência" in html
    assert f'value="{usuarios["Jaísa"].id}" selected' in html
    token = _token(html)

    resp = cli.post("/importacao/confirmar", data={"csrf_token": TOKEN_CSRF, "token": token, "responsavel_id": usuarios["Jaísa"].id,
                                                   "acao": "confirmar", "outras_carteiras": "1"},
                    follow_redirects=True)
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200
    assert "Importação #1 concluída" in html and "4 lançamento(s) novo(s)" in html
    assert "FATURAMENTO_V2.xlsx" in html and "Desfazer" in html
    imp = db.session.get(Importacao, 1)
    assert imp.responsavel.nome == "Jaísa" and imp.usuario.nome == "Natan"
    alfa = db.session.scalar(select(Empresa).where(Empresa.nome == "Loja Alfa (Matriz)"))
    assert alfa.historico_responsaveis[0].usuario.nome == "Jaísa"

    # token já usado não importa de novo
    resp = cli.post("/importacao/confirmar", data={"csrf_token": TOKEN_CSRF, "token": token, "responsavel_id": usuarios["Jaísa"].id},
                    follow_redirects=True)
    assert "expirada" in resp.get_data(as_text=True)
    assert db.session.query(Importacao).count() == 1

    resp = cli.post("/importacao/1/desfazer", data={"csrf_token": TOKEN_CSRF}, follow_redirects=True)
    assert "Importação #1 desfeita" in resp.get_data(as_text=True)
    db.session.expire_all()
    assert db.session.get(Importacao, 1).desfeita is True
    assert db.session.query(Lancamento).count() == 0


def test_tela_cancelar_e_erros(app, db, login, dados_tmp):
    cli = login("Natan")
    token = _token(_enviar(cli).get_data(as_text=True))
    resp = cli.post("/importacao/confirmar", data={"csrf_token": TOKEN_CSRF, "token": token, "acao": "cancelar"}, follow_redirects=True)
    assert "Nada foi gravado" in resp.get_data(as_text=True)
    assert not (dados_tmp / "importacoes" / f"{token}.json").exists()
    assert db.session.query(Importacao).count() == 0

    resp = _enviar(cli, b"lixo", nome="planilha.xlsx")
    assert resp.status_code == 302
    resp = cli.post("/importacao/previa", data={"csrf_token": TOKEN_CSRF, "arquivo": (io.BytesIO(b"a;b"), "x.csv")},
                    content_type="multipart/form-data", follow_redirects=True)
    assert "Formato não aceito" in resp.get_data(as_text=True)


def test_tela_exige_usuario(client):
    assert client.get("/importacao").status_code == 302


# ====================================================================== regressões da revisão

def _importar(db, usuarios, dados=None, quem="Natan", resp="Natan", **kw):
    imp = svc.confirmar(svc.analisar(dados or montar_planilha(), "f.xlsx"), usuarios[resp], usuarios[quem], **kw)
    db.session.commit()
    return imp


def test_desfazer_preserva_lancamento_criado_e_editado_depois(app, db, usuarios, tipos):
    from faturamento.services import lancamentos as svc_lanc
    imp = _importar(db, usuarios)
    alfa = db.session.scalar(select(Empresa).where(Empresa.nome == "Loja Alfa (Matriz)"))
    # depois da importação, Natan corrige os números de NF-e Entradas de 01/2026
    svc_lanc.salvar_mes(alfa, competencia("2026-01"), [{"tipo_id": tipos["NFE_ENT"].id, "qtde_autor": "11",
                                                         "valor_autor": "1300"}], usuarios["Natan"])
    db.session.commit()
    out = svc.desfazer(imp, usuarios["Natan"])
    db.session.commit()
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFE_ENT"])
    assert l is not None and (l.qtde_autor, l.valor_autor) == (11, Decimal("1300.00"))
    assert out["mantidos_editados"] == 1 and out["removidos"] == 3
    assert any("NF-e Entradas" in t for t in out["editados"])
    assert db.session.scalar(select(Empresa).where(Empresa.nome == "Loja Alfa (Matriz)")) is not None


def test_desfazer_preserva_lancamento_com_ajuste_de_outra_origem(app, db, usuarios, tipos):
    imp = _importar(db, usuarios)
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    aj = Ajuste(lancamento=l, grupo="AUTOR", qtde=1, valor=Decimal("5.00"), motivo="Outro", observacao="manual",
                criado_por_id=usuarios["Natan"].id)
    db.session.add(aj)  # ajuste gravado direto (sem mexer em alterado_em): só a origem do ajuste denuncia
    db.session.commit()
    out = svc.desfazer(imp, usuarios["Natan"])
    db.session.commit()
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert l is not None and out["mantidos_editados"] == 1
    assert sorted(a.motivo for a in l.ajustes) == ["Importado da planilha", "Outro"]  # nada dele foi removido


def test_desfazer_nao_restaura_lancamento_alterado_depois(app, db, usuarios, tipos, nova_empresa):
    from faturamento.services import lancamentos as svc_lanc
    alfa = nova_empresa("Loja Alfa (Matriz)", tipos_codigos=("NFCE",))
    db.session.add(Lancamento(empresa_id=alfa.id, competencia=competencia("2026-01"), tipo_id=tipos["NFCE"].id,
                              qtde_autor=1, valor_autor=Decimal("1.00")))
    db.session.commit()
    imp = _importar(db, usuarios)
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    svc_lanc.salvar_mes(alfa, competencia("2026-01"), [{"tipo_id": tipos["NFCE"].id, "qtde_autor": "101",
                                                         "valor_autor": "1001", "qtde_canc": "2", "valor_canc": "20"}],
                        usuarios["Natan"])
    db.session.commit()
    out = svc.desfazer(imp, usuarios["Natan"])
    db.session.commit()
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert (l.qtde_autor, l.valor_autor) == (101, Decimal("1001.00"))
    assert out["restaurados"] == 0 and out["mantidos_editados"] == 1
    assert [a.motivo for a in l.ajustes] == ["Importado da planilha"]  # ajuste da importação também fica


def test_desfazer_reimportacao_recria_ajuste_substituido(app, db, usuarios, tipos):
    _importar(db, usuarios)
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert [(a.qtde, a.valor) for a in l.ajustes] == [(5, Decimal("50.25"))]
    linhas = list(LINHAS_PADRAO)
    linhas[0] = _linha("LOJA ALFA (MATRIZ)", JAN, "NFC-e", 100, 1000.5, 7, 70, 2, 20)  # ajuste mudou na planilha
    imp2 = _importar(db, usuarios, montar_planilha(linhas))
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert [(a.qtde, a.valor) for a in l.ajustes] == [(7, Decimal("70.00"))]
    sub = svc.resumo(imp2)["ajustes_substituidos"]
    assert len(sub) == 1 and sub[0]["qtde"] == 5 and sub[0]["valor"] == "50.25"

    out = svc.desfazer(imp2, usuarios["Natan"])
    db.session.commit()
    db.session.expire_all()
    l = _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"])
    assert [(a.qtde, a.valor, a.motivo) for a in l.ajustes] == [(5, Decimal("50.25"), "Importado da planilha")]
    assert out["ajustes_recriados"] == 1 and out["ajustes_removidos"] == 1


def test_desfazer_so_quem_importou(app, db, usuarios):
    imp = _importar(db, usuarios)
    with pytest.raises(RegraViolada, match="Natan"):
        svc.desfazer(imp, usuarios["Hellen"])
    assert imp.desfeita is False


def test_confirmar_pula_carteira_de_outra_pessoa(app, db, usuarios, tipos, nova_empresa):
    nova_empresa("Loja Alfa (Matriz)", responsavel="Hellen", tipos_codigos=("NFCE",))
    previa = svc.analisar(montar_planilha(), "f.xlsx")
    assert previa.outras_carteiras(usuarios["Natan"], usuarios["Natan"]) == {"linhas": 2, "empresas": 1}
    imp = svc.confirmar(previa, usuarios["Natan"], usuarios["Natan"])
    db.session.commit()
    r = svc.resumo(imp)
    assert r["ignorados_outra_carteira"] == 2
    assert _lanc(db, "Loja Alfa (Matriz)", "2026-01", tipos["NFCE"]) is None
    assert any("carteira de Hellen" in a["texto"] for a in r["avisos"])
    assert r["lancamentos_criados"] == 2  # só a Padaria Beta (nova, do próprio Natan)


def test_confirmar_empresas_novas_para_outra_pessoa_exige_opcao(app, db, usuarios):
    previa = svc.analisar(montar_planilha(), "f.xlsx")
    with pytest.raises(RegraViolada, match="Importar também"):
        svc.confirmar(previa, usuarios["Mariana"], usuarios["Natan"])
    db.session.rollback()
    imp = svc.confirmar(previa, usuarios["Mariana"], usuarios["Natan"], outras_carteiras=True)
    assert svc.resumo(imp)["lancamentos_criados"] == 4


def test_tela_previa_mostra_outras_carteiras_e_desfazer_mostra_mantidos(app, db, login, usuarios, tipos, nova_empresa):
    from faturamento.services import lancamentos as svc_lanc
    nova_empresa("Loja Alfa (Matriz)", responsavel="Hellen", tipos_codigos=("NFCE",))
    cli = login("Natan")
    html = _enviar(cli, responsavel_id=usuarios["Natan"].id).get_data(as_text=True)
    assert "<strong>2</strong> lançamento(s) de 1 empresa(s) são de carteiras de outras pessoas" in html
    assert 'name="outras_carteiras"' in html and svc.OPCAO_OUTRAS in html
    resp = cli.post("/importacao/confirmar", data={"csrf_token": TOKEN_CSRF, "token": _token(html), "acao": "confirmar",
                                                   "responsavel_id": usuarios["Natan"].id}, follow_redirects=True)
    assert "2 lançamento(s) de carteiras de outras pessoas não foram importados" in resp.get_data(as_text=True)
    beta = db.session.scalar(select(Empresa).where(Empresa.nome == "Padaria Beta"))
    svc_lanc.salvar_mes(beta, competencia("2026-02"), [{"tipo_id": tipos["NFSE_SAI"].id, "qtde_autor": "4",
                                                         "valor_autor": "400"}], usuarios["Natan"])
    db.session.commit()
    resp = cli.post("/importacao/1/desfazer", data={"csrf_token": TOKEN_CSRF}, follow_redirects=True)
    html = resp.get_data(as_text=True)
    assert "1 lançamento(s) alterado(s) depois da importação foram mantidos" in html
    assert "Padaria Beta · 02/2026 · NFS-e Saídas" in html


def test_avaliar_conta_formula_grande_ou_profunda():
    assert svc.avaliar_conta("=" + "-" * 5000 + "1") is None
    assert svc.avaliar_conta("=" + "1+" * 300 + "1") is None   # > 500 caracteres
    assert svc.avaliar_conta("=" + "(" * 240 + "1" + ")" * 240) is None
    assert svc.avaliar_conta("=1+2") == Decimal("3")


def test_formula_gigante_vira_aviso_na_previa(app, db):
    dados = montar_planilha([_linha("LOJA X", JAN, "NFC-e", 1, "=" + "-" * 3000 + "1")])
    p = svc.analisar(dados, "f.xlsx")
    assert any("não consegui calcular" in a.texto and a.destaque for a in p.avisos)

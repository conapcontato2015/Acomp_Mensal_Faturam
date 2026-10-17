"""Conferência (cartões por empresa) e a planilha Excel exportada."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from faturamento.db import db as _db
from faturamento.services import exportacao as svc_exp
from faturamento.services import lancamentos as svc_lanc
from faturamento.services import resumo as svc_res


def _lancar(empresa, comp, usuario, tipos, valores):
    """valores: {codigo: (qtde_autor, valor_autor, qtde_canc, valor_canc)}"""
    linhas = [{"tipo_id": tipos[c].id, "qtde_autor": v[0], "valor_autor": v[1], "qtde_canc": v[2], "valor_canc": v[3]}
              for c, v in valores.items()]
    svc_lanc.salvar_mes(empresa, comp, linhas, usuario)
    _db.session.commit()
    return {c: svc_lanc.buscar(empresa.id, comp, tipos[c].id) for c in valores}


@pytest.fixture()
def cenario(app, nova_empresa, usuarios, tipos, comp):
    natan, hellen = usuarios["Natan"], usuarios["Hellen"]
    alfa = nova_empresa("Alfa Comércio", "Natan")
    beta = nova_empresa("Beta Serviços", "Natan")
    gama = nova_empresa("Gama Hellen", "Hellen")
    la = _lancar(alfa, comp, natan, tipos, {
        "NFE_ENT": (168, "190505.28", 1, "328.99"),
        "NFE_SAI": (18, "9593.45", None, None),
        "NFCE": (6431, "186864.95", 27, "727.40"),
    })
    svc_lanc.adicionar_ajuste(la["NFCE"], "AUTOR", 61, "1295.50", natan)
    _db.session.commit()
    lb = _lancar(beta, comp, natan, tipos, {"NFE_ENT": (3, "22.43", None, None)})
    svc_lanc.marcar_bateu(lb["NFE_ENT"], True, natan)
    lg = _lancar(gama, comp, hellen, tipos, {"NFE_SAI": (5, "1000.00", None, None)})
    _db.session.commit()
    return {"alfa": alfa, "beta": beta, "gama": gama, "la": la, "lb": lb, "lg": lg}


def _ler_xlsx(conteudo: bytes):
    wb = load_workbook(BytesIO(conteudo))
    ws = wb.active
    return wb, ws, [list(r) for r in ws.iter_rows(values_only=True)]


def test_cartoes_da_carteira_e_de_todos(cenario, usuarios, comp):
    natan = usuarios["Natan"]
    meus = svc_res.conferencia(comp, natan.id)
    assert [c.empresa.nome for c in meus] == ["Alfa Comércio", "Beta Serviços"]
    alfa, beta = meus
    assert [b.sentido for b in alfa.blocos] == ["E", "S"]  # entradas e saídas em blocos separados
    nfce = next(l for b in alfa.blocos for l in b.linhas if l.tipo.codigo == "NFCE")
    assert nfce.total_qtde_autor == 6431 + 61  # já com o ajuste
    assert (alfa.total, alfa.n_bateu, alfa.tudo_bateu) == (3, 0, False)
    assert (beta.total, beta.n_bateu, beta.tudo_bateu) == (1, 1, True)
    assert len(svc_res.conferencia(comp, None)) == 3


def test_excel(cenario, usuarios, comp):
    assert svc_exp.nome_arquivo_conferencia(comp) == "conferencia-2026-08.xlsx"
    wb, ws, linhas = _ler_xlsx(svc_exp.planilha_conferencia(comp, svc_res.conferencia(comp, usuarios["Natan"].id)))
    assert ws.title == "Conferência 08-2026"
    assert linhas[0] == ["Empresa", "Responsável", "Sentido", "Tipo", "Qtde autorizada", "Valor autorizado",
                         "Qtde cancelada", "Valor cancelado", "Ajuste qtde", "Ajuste valor", "Bateu"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold
    corpo = linhas[1:]
    # Alfa: 1 entrada + subtotal, 2 saídas + subtotal; Beta: 1 entrada + subtotal
    assert len(corpo) == 7
    nfce = next(l for l in corpo if l[0] == "Alfa Comércio" and l[3] == "NFC-e")
    assert nfce[4] == 6492 and Decimal(str(nfce[5])) == Decimal("188160.45")
    assert nfce[8] == 61 and Decimal(str(nfce[9])) == Decimal("1295.50") and nfce[10] == "Não"
    beta = next(l for l in corpo if l[0] == "Beta Serviços" and l[3] == "NF-e Entradas")
    assert beta[10] == "Sim"

    subs = [(i, l) for i, l in enumerate(corpo, start=2) if str(l[3]).startswith("Subtotal")]
    assert [(l[0], l[2], l[3]) for _, l in subs] == [
        ("Alfa Comércio", "Entrada", "Subtotal entradas"),
        ("Alfa Comércio", "Saída", "Subtotal saídas"),
        ("Beta Serviços", "Entrada", "Subtotal entradas"),
    ]
    sub_e, sub_s = subs[0][1], subs[1][1]
    assert sub_e[4] == 168 and Decimal(str(sub_e[5])) == Decimal("190505.28")
    assert sub_s[4] == 18 + 6492 and Decimal(str(sub_s[5])) == Decimal("9593.45") + Decimal("188160.45")
    # entradas e saídas nunca somadas: nenhum valor da planilha é E+S da Alfa
    soma_es = Decimal("190505.28") + Decimal("9593.45") + Decimal("188160.45")
    assert all(Decimal(str(l[5])) != soma_es for l in corpo if l[5] is not None)
    assert not any("total" in str(l[3]).lower() and "subtotal" not in str(l[3]).lower() for l in corpo)
    for i, _ in subs:
        assert ws.cell(row=i, column=1).font.bold
    assert ws.cell(row=2, column=5).number_format == "#,##0"
    assert ws.cell(row=2, column=6).number_format == "#,##0.00"


def test_excel_carteira_vazia(app, comp):
    _, _, linhas = _ler_xlsx(svc_exp.planilha_conferencia(comp, []))
    assert len(linhas) == 1

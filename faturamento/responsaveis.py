"""De onde vem o responsável de cada empresa.

Hoje a carteira é mantida à mão (qualquer usuário troca) ou por planilha. Quando o sistema
externo que define os responsáveis for ligado, basta implementar ``ProvedorApi.buscar`` (ou
escrever outro provedor) e trocar ``RESPONSAVEIS_PROVEDOR`` no .env — o resto do sistema não muda.

Todo provedor devolve uma lista de ``Atribuicao``; ``sincronizar`` aplica a lista gravando a
origem (planilha/integração) no histórico e devolve um relatório.
"""
from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import ArgumentError, SQLAlchemyError
from sqlalchemy.pool import NullPool

from .config import config
from .extensions import db
from .models import ORIGEM_INTEGRACAO, ORIGEM_PLANILHA, Usuario
from .services.empresas import buscar_por_cnpj, buscar_por_nome, definir_responsavel, responsavel_em
from .util import ValorInvalido, competencia, competencia_atual, normalizar_nome, so_digitos


@dataclass
class Atribuicao:
    responsavel: str
    cnpj: str | None = None
    empresa: str | None = None
    inicio: date | None = None
    linha: int | None = None
    # Dados extras que alguns provedores trazem (ex.: Controle Fiscal)
    codigo: str | None = None
    ativo: bool | None = None


@dataclass
class RelatorioSincronizacao:
    alteradas: list[str] = field(default_factory=list)
    sem_mudanca: int = 0
    empresas_nao_encontradas: list[str] = field(default_factory=list)
    responsaveis_nao_encontrados: list[str] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)

    def como_dict(self) -> dict:
        return {
            "alteradas": self.alteradas, "sem_mudanca": self.sem_mudanca,
            "empresas_nao_encontradas": self.empresas_nao_encontradas,
            "responsaveis_nao_encontrados": self.responsaveis_nao_encontrados, "erros": self.erros,
        }


class ProvedorResponsaveis:
    origem = ORIGEM_INTEGRACAO
    nome = "base"

    def buscar(self) -> list[Atribuicao]:  # pragma: no cover - interface
        raise NotImplementedError


class ProvedorManual(ProvedorResponsaveis):
    """Padrão: sem fonte externa. A carteira é editada na tela de Empresas."""

    nome = "manual"

    def buscar(self) -> list[Atribuicao]:
        return []


class ProvedorPlanilha(ProvedorResponsaveis):
    """Lê .xlsx ou .csv com colunas: CNPJ e/ou EMPRESA, RESPONSAVEL e (opcional) INICIO (AAAA-MM ou MM/AAAA)."""

    origem = ORIGEM_PLANILHA
    nome = "planilha"

    def __init__(self, conteudo: bytes, nome_arquivo: str):
        self.conteudo = conteudo
        self.nome_arquivo = nome_arquivo

    def _linhas(self) -> list[list]:
        try:
            if self.nome_arquivo.lower().endswith((".xlsx", ".xlsm")):
                wb = load_workbook(io.BytesIO(self.conteudo), read_only=True, data_only=True)
                try:
                    ws = wb.worksheets[0]
                    return [list(r) for r in ws.iter_rows(values_only=True)]
                finally:
                    wb.close()
            texto = self.conteudo.decode("utf-8-sig", errors="replace")
            if not texto.strip():
                return []
            try:
                dialeto = csv.Sniffer().sniff(texto[:2048], delimiters=";,\t")
            except csv.Error:
                # Sniffer falha em arquivo de uma coluna só: usa o padrão do Excel brasileiro
                dialeto = type("PontoEVirgula", (csv.excel,), {"delimiter": ";"})
            return [r for r in csv.reader(io.StringIO(texto), dialeto)]
        except (csv.Error, zipfile.BadZipFile, InvalidFileException, KeyError, UnicodeDecodeError,
                IndexError, OSError) as exc:
            raise ValorInvalido(f"Arquivo ilegível: {exc}") from exc

    def buscar(self) -> list[Atribuicao]:
        linhas = self._linhas()
        if not linhas:
            return []
        cab = [normalizar_nome(c) for c in linhas[0]]

        def idx(*nomes):
            for n in nomes:
                if n in cab:
                    return cab.index(n)
            return None

        i_cnpj, i_emp = idx("CNPJ"), idx("EMPRESA", "RAZAO SOCIAL", "NOME")
        i_resp, i_ini = idx("RESPONSAVEL", "RESPONSÁVEL", "USUARIO"), idx("INICIO", "A PARTIR DE", "COMPETENCIA")
        if i_resp is None or (i_cnpj is None and i_emp is None):
            raise ValorInvalido("A planilha precisa das colunas RESPONSAVEL e CNPJ ou EMPRESA (na primeira linha).")
        saida = []
        for n, linha in enumerate(linhas[1:], start=2):
            def val(i):
                return linha[i] if i is not None and i < len(linha) and linha[i] is not None else None
            resp = val(i_resp)
            if not resp or not str(resp).strip():
                continue
            ini = val(i_ini)
            saida.append(Atribuicao(
                responsavel=str(resp).strip(),
                cnpj=so_digitos(val(i_cnpj)) or None,
                empresa=str(val(i_emp)).strip() if val(i_emp) else None,
                inicio=competencia(ini) if ini else None,
                linha=n,
            ))
        return saida


class ProvedorApi(ProvedorResponsaveis):
    """Ponto de encaixe do sistema externo de responsáveis (ainda não definido).

    Para ligar: implemente ``buscar`` chamando a API/arquivo/banco do outro sistema e
    devolvendo ``Atribuicao(responsavel=..., cnpj=..., inicio=...)``. Configure no .env:
    RESPONSAVEIS_PROVEDOR=api, RESPONSAVEIS_API_URL=..., RESPONSAVEIS_API_TOKEN=...
    """

    nome = "api"

    def __init__(self, url: str, token: str = ""):
        self.url = url
        self.token = token

    def buscar(self) -> list[Atribuicao]:
        raise NotImplementedError(
            "Integração com o sistema de responsáveis ainda não implementada. "
            "Enquanto isso, troque na tela de Empresas ou importe uma planilha."
        )


def ler_mapa_responsaveis(texto: str | None) -> dict[str, str]:
    """'NomeLa=NomeAqui;Outro=Natan' -> {'NOMELA': 'NomeAqui', 'OUTRO': 'Natan'} (chave normalizada)."""
    mapa = {}
    for par in (texto or "").replace("\n", ";").split(";"):
        la, sep, aqui = par.partition("=")
        if sep and la.strip() and aqui.strip():
            mapa[normalizar_nome(la)] = aqui.strip()
    return mapa


class ProvedorControleFiscal(ProvedorResponsaveis):
    """Lê empresas e responsáveis direto do banco do sistema Controle Fiscal — SOMENTE LEITURA.

    Os dois sistemas rodam no mesmo PC: o Faturamento abre uma conexão própria ao PostgreSQL do
    Controle Fiscal (FOLHA_DATABASE_URL), com a sessão travada em modo só leitura, e faz apenas
    SELECT. Nomes de responsável diferentes entre os sistemas são traduzidos por
    FOLHA_MAPA_RESPONSAVEIS ("NomeLa=NomeAqui;...").
    """

    nome = "controle_fiscal"
    rotulo = "Controle Fiscal"
    TIMEOUT_CONEXAO = 5  # segundos

    CONSULTA = text(
        "SELECT e.codigo, e.razao_social, e.cnpj, e.ativo, r.nome AS responsavel "
        "FROM empresa e LEFT JOIN responsavel r ON r.id = e.responsavel_id "
        "ORDER BY e.razao_social"
    )

    def __init__(self, url: str, mapa: dict[str, str] | str | None = None):
        self.url = (url or "").strip()
        self.mapa = ler_mapa_responsaveis(mapa) if isinstance(mapa, str) or mapa is None else mapa

    @classmethod
    def do_config(cls) -> "ProvedorControleFiscal | None":
        url = config.get("FOLHA_DATABASE_URL", "")
        if not url:
            return None
        return cls(url, config.get("FOLHA_MAPA_RESPONSAVEIS", ""))

    def _engine(self):
        url = make_url(self.url)
        opcoes: dict = {"poolclass": NullPool}
        if url.get_backend_name() == "postgresql":
            # A sessão inteira nasce só leitura (equivale a SET SESSION CHARACTERISTICS AS TRANSACTION
            # READ ONLY): qualquer INSERT/UPDATE/DELETE seria recusado pelo próprio PostgreSQL.
            opcoes["connect_args"] = {"connect_timeout": self.TIMEOUT_CONEXAO,
                                      "options": "-c default_transaction_read_only=on"}
        elif url.get_backend_name() == "sqlite" and url.database and url.database != ":memory:":
            # Usado nos testes: abre o arquivo em modo só leitura.
            url = url.set(database=f"file:{url.database}", query={"mode": "ro", "uri": "true"})
        return create_engine(url, **opcoes)

    def _mapear(self, nome: str | None) -> str:
        nome = (nome or "").strip()
        return self.mapa.get(normalizar_nome(nome), nome) if nome else ""

    def buscar(self) -> list[Atribuicao]:
        try:
            engine = self._engine()
        except (ArgumentError, ValueError) as exc:
            raise ValorInvalido(
                "Não foi possível conectar ao banco do Controle Fiscal: endereço inválido. "
                "Verifique se o PostgreSQL dele está ligado e o FOLHA_DATABASE_URL no .env.") from exc
        except ImportError as exc:
            raise ValorInvalido(
                f"Não foi possível conectar ao banco do Controle Fiscal: driver ausente ({exc}). "
                "Rode: pip install -r requirements.txt") from exc
        try:
            with engine.connect() as conn:
                linhas = conn.execute(self.CONSULTA).mappings().all()
                conn.rollback()
        except (SQLAlchemyError, ImportError) as exc:
            detalhe = str(getattr(exc, "orig", None) or exc).strip().splitlines()
            raise ValorInvalido(
                f"Não foi possível conectar ao banco do Controle Fiscal: {detalhe[0] if detalhe else exc.__class__.__name__}. "
                "Verifique se o PostgreSQL dele está ligado e o FOLHA_DATABASE_URL no .env.") from exc
        finally:
            engine.dispose()
        saida = []
        for ln in linhas:
            ativo = ln["ativo"]
            saida.append(Atribuicao(
                responsavel=self._mapear(ln["responsavel"]),
                cnpj=so_digitos(ln["cnpj"]) or None,
                empresa=(ln["razao_social"] or "").strip() or None,
                codigo=str(ln["codigo"]).strip() if ln["codigo"] is not None else None,
                ativo=None if ativo is None else bool(ativo),
            ))
        return saida


def provedor_configurado() -> ProvedorResponsaveis:
    tipo = config.get("RESPONSAVEIS_PROVEDOR", "manual")
    if tipo == "controle_fiscal":
        return ProvedorControleFiscal(config.get("FOLHA_DATABASE_URL", ""),
                                      config.get("FOLHA_MAPA_RESPONSAVEIS", ""))
    if tipo == "api":
        return ProvedorApi(config.get("RESPONSAVEIS_API_URL", ""),
                           config.get("RESPONSAVEIS_API_TOKEN", ""))
    return ProvedorManual()


def _achar_usuario(nome: str) -> Usuario | None:
    alvo = normalizar_nome(nome)
    for u in db.session.query(Usuario).all():
        if normalizar_nome(u.nome) == alvo or (u.external_id and u.external_id == nome):
            return u
    return None


def sincronizar(provedor: ProvedorResponsaveis, registrado_por: Usuario | None,
                inicio_padrao: date | None = None, aplicar: bool = True) -> RelatorioSincronizacao:
    """Aplica as atribuições do provedor. Com ``aplicar=False`` só simula (prévia)."""
    rel = RelatorioSincronizacao()
    inicio_padrao = inicio_padrao or competencia_atual()
    try:
        atribs = provedor.buscar()
    except (ValorInvalido, NotImplementedError) as exc:
        rel.erros.append(str(exc))
        return rel
    for a in atribs:
        if not (a.responsavel or "").strip():
            continue  # fonte sem responsável para esta empresa: nada a aplicar
        empresa = buscar_por_cnpj(a.cnpj) if a.cnpj else None
        if empresa is None and a.empresa:
            empresa = buscar_por_nome(a.empresa)
        rotulo = a.empresa or a.cnpj or f"linha {a.linha}"
        if empresa is None:
            rel.empresas_nao_encontradas.append(rotulo)
            continue
        usuario = _achar_usuario(a.responsavel)
        if usuario is None:
            rel.responsaveis_nao_encontrados.append(f"{a.responsavel} ({rotulo})")
            continue
        inicio = a.inicio or inicio_padrao
        atual = responsavel_em(empresa, inicio)
        if atual is not None and atual.id == usuario.id:
            rel.sem_mudanca += 1
            continue
        rel.alteradas.append(f"{empresa.nome}: {atual.nome if atual else '—'} → {usuario.nome}")
        if aplicar:
            definir_responsavel(empresa, usuario, inicio, registrado_por, origem=provedor.origem)
    return rel

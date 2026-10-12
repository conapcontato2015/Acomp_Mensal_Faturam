"""Importação da planilha antiga (FATURAMENTO_V2.xlsx, aba DADOS).

Fluxo em duas etapas:

1. ``analisar(bytes, nome)`` lê a planilha e devolve uma ``Previa`` (nada é gravado).
   A prévia é guardada em ``DADOS_DIR/importacoes/<token>.json`` entre as requisições.
2. ``confirmar(previa, responsavel, usuario, nome)`` grava tudo numa transação e devolve a
   ``Importacao``. ``desfazer(importacao, usuario)`` remove o que ela criou.

A planilha é lida pelos NOMES das colunas (linha de cabeçalho com "EMPRESA" nas 10 primeiras
linhas), não por letras fixas. Totais (fórmulas) são ignorados: o sistema recalcula.
"""
from __future__ import annotations

import ast
import io
import json
import operator
import re
import secrets
import time
import zipfile
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from flask import current_app
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select

from ..extensions import db
from ..models import (
    GRUPO_AUTOR, GRUPO_CANC, ORIGEM_PLANILHA, Ajuste, Empresa, Fechamento, Importacao, Lancamento,
    TipoDocumento, Usuario, agora,
)
from ..util import (
    ValorInvalido, comp_curta, comp_iso, competencia, competencia_atual, fmt_inteiro, fmt_moeda,
    nome_exibicao, normalizar_nome, parse_decimal, parse_inteiro,
)
from . import RegraViolada, auditoria
from .empresas import atualizar_empresa, criar_empresa, responsavel_em
from .lancamentos import CAMPOS, limpar_sem_movimento

MOTIVO_IMPORTACAO = "Importado da planilha"

# nome normalizado da coluna -> campo interno
COLUNAS = {
    "EMPRESA": "empresa",
    "TIPO": "tipo",
    "QTDE_AUTOR": "qtde_autor",
    "VALOR_AUTOR": "valor_autor",
    "AJUSTE_QTDE_AUTOR": "aj_qtde_autor",
    "AJUSTE_VALOR_AUTOR": "aj_valor_autor",
    "QTDE_CANC": "qtde_canc",
    "VALOR_CANC": "valor_canc",
    "AJUSTE_QTDE_CANC": "aj_qtde_canc",
    "AJUSTE_VALOR_CANC": "aj_valor_canc",
}
CAMPOS_ENTRADA = ("qtde_autor", "valor_autor", "aj_qtde_autor", "aj_valor_autor",
                  "qtde_canc", "valor_canc", "aj_qtde_canc", "aj_valor_canc")
OBRIGATORIAS = ("empresa", "comp", "tipo", "qtde_autor", "valor_autor")
ROTULO_COLUNA = {v: k for k, v in COLUNAS.items()} | {"comp": "COMP."}

_TOKEN_OK = re.compile(r"^[A-Za-z0-9_-]{10,64}$")


# ====================================================================== prévia

@dataclass
class Item:
    """Uma linha útil da planilha (empresa, competência, tipo). Decimais como texto."""

    linha: int
    empresa: str           # como está na planilha
    comp: str              # AAAA-MM
    tipo_codigo: str
    tipo_nome: str
    qtde_autor: int | None = None
    valor_autor: str | None = None
    qtde_canc: int | None = None
    valor_canc: str | None = None
    aj_qtde_autor: int = 0
    aj_valor_autor: str = "0.00"
    aj_qtde_canc: int = 0
    aj_valor_canc: str = "0.00"

    @property
    def competencia(self) -> date:
        return competencia(self.comp)

    def base(self) -> dict:
        return {
            "qtde_autor": self.qtde_autor,
            "valor_autor": Decimal(self.valor_autor) if self.valor_autor is not None else None,
            "qtde_canc": self.qtde_canc,
            "valor_canc": Decimal(self.valor_canc) if self.valor_canc is not None else None,
        }

    def ajuste(self, grupo: str) -> tuple[int, Decimal]:
        if grupo == GRUPO_AUTOR:
            return self.aj_qtde_autor, Decimal(self.aj_valor_autor)
        return self.aj_qtde_canc, Decimal(self.aj_valor_canc)

    @property
    def n_ajustes(self) -> int:
        return sum(1 for g in (GRUPO_AUTOR, GRUPO_CANC) if any(self.ajuste(g)))


@dataclass
class Aviso:
    texto: str
    linha: int | None = None
    destaque: bool = False   # precisa de correção manual depois da importação


@dataclass
class Previa:
    arquivo: str
    aba: str = ""
    linha_cabecalho: int = 0
    linhas_lidas: int = 0
    linhas_vazias: int = 0
    itens: list[Item] = field(default_factory=list)
    avisos: list[Aviso] = field(default_factory=list)
    token: str | None = None

    # ---------------------------------------------------------------- serialização
    def como_dict(self) -> dict:
        d = asdict(self)
        return d

    @classmethod
    def de_dict(cls, d: dict) -> "Previa":
        p = cls(arquivo=d["arquivo"], aba=d.get("aba", ""), linha_cabecalho=d.get("linha_cabecalho", 0),
                linhas_lidas=d.get("linhas_lidas", 0), linhas_vazias=d.get("linhas_vazias", 0),
                token=d.get("token"))
        p.itens = [Item(**i) for i in d.get("itens", [])]
        p.avisos = [Aviso(**a) for a in d.get("avisos", [])]
        return p

    # ---------------------------------------------------------------- estatísticas
    def avisar(self, texto: str, linha: int | None = None, destaque: bool = False) -> None:
        self.avisos.append(Aviso(texto, linha, destaque))

    @property
    def avisos_destaque(self) -> list[Aviso]:
        return [a for a in self.avisos if a.destaque]

    @property
    def competencias(self) -> list[date]:
        return sorted({i.competencia for i in self.itens})

    @property
    def n_lancamentos(self) -> int:
        return len(self.itens)

    @property
    def n_ajustes(self) -> int:
        return sum(i.n_ajustes for i in self.itens)

    def empresas(self) -> list[dict]:
        """Uma linha por empresa da planilha, indicando se já existe no sistema."""
        existentes = _empresas_por_nome()
        grupos: dict[str, list[Item]] = defaultdict(list)
        originais: dict[str, str] = {}
        for i in self.itens:
            chave = normalizar_nome(i.empresa)
            grupos[chave].append(i)
            originais.setdefault(chave, i.empresa)
        saida = []
        for chave, itens in grupos.items():
            emp = existentes.get(chave)
            comps = sorted({i.competencia for i in itens})
            tipos = {i.tipo_codigo: i.tipo_nome for i in itens}
            saida.append({
                "original": originais[chave],
                "nome": emp.nome if emp else nome_exibicao(originais[chave]),
                "nova": emp is None,
                "empresa_id": emp.id if emp else None,
                "competencias": comps,
                "primeira": comps[0], "ultima": comps[-1],
                "n_lancamentos": len(itens),
                "n_ajustes": sum(i.n_ajustes for i in itens),
                "tipos": [tipos[c] for c in _ordem_tipos() if c in tipos],
            })
        saida.sort(key=lambda e: normalizar_nome(e["nome"]))
        return saida

    def totais_por_competencia(self) -> list[dict]:
        """Somas com ajuste, separadas por sentido (entradas e saídas nunca se somam)."""
        sentido = {t.codigo: t.sentido for t in db.session.scalars(select(TipoDocumento))}
        acc: dict[date, dict] = {}
        for i in self.itens:
            c = i.competencia
            t = acc.setdefault(c, {
                "competencia": c, "lancamentos": 0, "empresas": set(),
                "E": {"qtde": 0, "valor": Decimal("0"), "qtde_canc": 0, "valor_canc": Decimal("0")},
                "S": {"qtde": 0, "valor": Decimal("0"), "qtde_canc": 0, "valor_canc": Decimal("0")},
            })
            t["lancamentos"] += 1
            t["empresas"].add(normalizar_nome(i.empresa))
            b = i.base()
            aq, av = i.ajuste(GRUPO_AUTOR)
            cq, cv = i.ajuste(GRUPO_CANC)
            s = t[sentido.get(i.tipo_codigo, "S")]
            s["qtde"] += (b["qtde_autor"] or 0) + aq
            s["valor"] += (b["valor_autor"] or 0) + av
            s["qtde_canc"] += (b["qtde_canc"] or 0) + cq
            s["valor_canc"] += (b["valor_canc"] or 0) + cv
        saida = []
        for c in sorted(acc):
            t = acc[c]
            t["empresas"] = len(t["empresas"])
            saida.append(t)
        return saida

    def conflitos(self) -> dict:
        """O que já existe no banco para as mesmas chaves (para avisar na prévia)."""
        existentes = _empresas_por_nome()
        tipos = {t.codigo: t.id for t in db.session.scalars(select(TipoDocumento))}
        ids = {e.id for k, e in existentes.items() if any(normalizar_nome(i.empresa) == k for i in self.itens)}
        atuais = _lancamentos_existentes(ids)
        iguais = alterar = bateu = 0
        for i in self.itens:
            emp = existentes.get(normalizar_nome(i.empresa))
            if emp is None:
                continue
            l = atuais.get((emp.id, i.competencia, tipos.get(i.tipo_codigo)))
            if l is None:
                continue
            if l.bateu:
                bateu += 1
            elif _mesmos_numeros(l, i.base()):
                iguais += 1
            else:
                alterar += 1
        return {"iguais": iguais, "alterar": alterar, "bateu": bateu}


def _ordem_tipos() -> list[str]:
    return [t.codigo for t in db.session.scalars(select(TipoDocumento).order_by(TipoDocumento.ordem))]


def _empresas_por_nome() -> dict[str, Empresa]:
    return {normalizar_nome(e.nome): e for e in db.session.scalars(select(Empresa))}


def _lancamentos_existentes(empresa_ids) -> dict[tuple, Lancamento]:
    if not empresa_ids:
        return {}
    q = select(Lancamento).where(Lancamento.empresa_id.in_(list(empresa_ids)))
    return {(l.empresa_id, l.competencia, l.tipo_id): l for l in db.session.scalars(q)}


def _mesmos_numeros(lanc: Lancamento, base: dict) -> bool:
    for c in CAMPOS:
        x, y = getattr(lanc, c), base.get(c)
        if x is None or y is None:
            if not (x is None and y is None):
                return False
        elif Decimal(x) != Decimal(y):
            return False
    return True


# ====================================================================== leitura

_OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}


def avaliar_conta(formula: str) -> Decimal | None:
    """Avalia contas simples digitadas na célula ('=75536.37-11.5'). Só números e + - * / ( )."""
    texto = (formula or "").strip().lstrip("=").replace(",", ".")
    if not texto or not re.fullmatch(r"[\d.+\-*/() ]+", texto):
        return None
    try:
        arvore = ast.parse(texto, mode="eval")
    except SyntaxError:
        return None

    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return Decimal(str(n.value))
        if isinstance(n, ast.BinOp) and type(n.op) in _OPS:
            return _OPS[type(n.op)](ev(n.left), ev(n.right))
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.USub, ast.UAdd)):
            v = ev(n.operand)
            return -v if isinstance(n.op, ast.USub) else v
        raise ValueError("expressão não suportada")

    try:
        return ev(arvore)
    except (ValueError, ArithmeticError):
        return None


def _abrir(conteudo: bytes, data_only: bool):
    try:
        return load_workbook(io.BytesIO(conteudo), data_only=data_only, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValorInvalido("O arquivo não é uma planilha .xlsx válida.") from exc


def _achar_aba_e_cabecalho(wb) -> tuple[str, int, dict[str, int]]:
    nomes = list(wb.sheetnames)
    ordem = [n for n in nomes if normalizar_nome(n) == "DADOS"] + [n for n in nomes if normalizar_nome(n) != "DADOS"]
    for nome in ordem:
        ws = wb[nome]
        for n, linha in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            cab = [normalizar_nome(v) if v is not None else "" for v in linha]
            if "EMPRESA" not in cab:
                continue
            mapa: dict[str, int] = {}
            for idx, c in enumerate(cab):
                if c in COLUNAS and COLUNAS[c] not in mapa:
                    mapa[COLUNAS[c]] = idx
                elif (c.startswith("COMP") or c == "MES") and "comp" not in mapa:
                    mapa["comp"] = idx
            faltam = [ROTULO_COLUNA[c] for c in OBRIGATORIAS if c not in mapa]
            if faltam:
                raise ValorInvalido(f"Aba {nome}: faltam as colunas {', '.join(faltam)} no cabeçalho (linha {n}).")
            return nome, n, mapa
    raise ValorInvalido("Não encontrei a linha de cabeçalho (com a coluna EMPRESA) nas 10 primeiras linhas. "
                        "Use a aba DADOS da planilha FATURAMENTO_V2.")


def _ler_competencia(v) -> date:
    if isinstance(v, (datetime, date)):
        return competencia(v)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return competencia(from_excel(v))
    return competencia(str(v).strip())


def analisar(conteudo: bytes, nome_arquivo: str) -> Previa:
    """Lê a planilha e monta a prévia. Não grava nada no banco."""
    if not conteudo:
        raise ValorInvalido("Arquivo vazio.")
    wb_val = _abrir(conteudo, data_only=True)
    wb_form = _abrir(conteudo, data_only=False)
    try:
        aba, n_cab, mapa = _achar_aba_e_cabecalho(wb_val)
        previa = Previa(arquivo=nome_arquivo, aba=aba, linha_cabecalho=n_cab)
        tipos = {normalizar_nome(t.nome): t for t in db.session.scalars(select(TipoDocumento))}
        tipos.update({normalizar_nome(t.codigo): t for t in tipos.copy().values()})
        hoje = competencia_atual()
        vistos: dict[tuple, int] = {}
        tipos_desconhecidos: dict[str, list[int]] = defaultdict(list)
        futuras: dict[str, list[int]] = defaultdict(list)

        linhas_val = wb_val[aba].iter_rows(min_row=n_cab + 1, values_only=True)
        linhas_form = wb_form[aba].iter_rows(min_row=n_cab + 1, values_only=True)
        for n, (lv, lf) in enumerate(zip(linhas_val, linhas_form), start=n_cab + 1):
            previa.linhas_lidas += 1

            def cel(campo, linha=lv):
                i = mapa.get(campo)
                return linha[i] if i is not None and i < len(linha) else None

            brutos: dict[str, object] = {}
            for campo in CAMPOS_ENTRADA:
                if campo not in mapa:
                    continue
                v = cel(campo)
                f = cel(campo, lf)
                if isinstance(f, str) and f.startswith("="):
                    ref = f"{get_column_letter(mapa[campo] + 1)}{n}"
                    if v is None:
                        v = avaliar_conta(f)
                    if v is None:
                        previa.avisar(f"{ref}: célula com fórmula que não consegui calcular ({f}); deixada em branco — "
                                      "confira no lançamento", n, destaque=True)
                    else:
                        try:
                            num = fmt_inteiro(parse_inteiro(v)) if "qtde" in campo else fmt_moeda(parse_decimal(v))
                        except ValorInvalido:
                            num = str(v)
                        previa.avisar(f"{ref}: célula com conta digitada ({f}); importado o resultado ({num})", n)
                if isinstance(v, str) and not v.strip():
                    v = None
                brutos[campo] = v

            if all(v is None for v in brutos.values()):
                previa.linhas_vazias += 1
                continue

            emp = cel("empresa")
            emp = " ".join(str(emp).split()) if emp is not None else ""
            if not emp:
                previa.avisar(f"Linha {n}: tem números mas está sem EMPRESA — ignorada", n, destaque=True)
                continue
            try:
                comp = _ler_competencia(cel("comp"))
            except (ValorInvalido, ValueError, TypeError, OverflowError):
                previa.avisar(f"Linha {n}: competência inválida ({cel('comp')!r}) em {emp} — ignorada", n, destaque=True)
                continue
            tipo_txt = cel("tipo")
            tipo = tipos.get(normalizar_nome(tipo_txt)) if tipo_txt is not None else None
            if tipo is None:
                tipos_desconhecidos[str(tipo_txt or "(vazio)")].append(n)
                continue
            if comp > hoje:
                futuras[comp_curta(comp)].append(n)
                continue
            rotulo = f"{tipo.nome}, {nome_exibicao(emp)}, {comp_curta(comp)}"

            # números
            valores: dict[str, object] = {}
            ruim = False
            for campo, v in brutos.items():
                try:
                    valores[campo] = parse_inteiro(v) if "qtde" in campo else parse_decimal(v)
                except ValorInvalido:
                    previa.avisar(f"Linha {n}: '{v}' em {ROTULO_COLUNA[campo]} não é número ({rotulo}) — linha ignorada",
                                  n, destaque=True)
                    ruim = True
            if ruim:
                continue
            negativos = [ROTULO_COLUNA[c] for c in ("qtde_autor", "valor_autor", "qtde_canc", "valor_canc")
                         if valores.get(c) is not None and valores[c] < 0]
            if negativos:
                previa.avisar(f"Linha {n}: número negativo em {', '.join(negativos)} ({rotulo}) — linha ignorada; "
                              "diferenças entram como ajuste", n, destaque=True)
                continue

            chave = (normalizar_nome(emp), comp, tipo.codigo)
            if chave in vistos:
                previa.avisar(f"Linha {n}: repete empresa, competência e tipo da linha {vistos[chave]} ({rotulo}) — "
                              "mantida a primeira", n, destaque=True)
                continue
            vistos[chave] = n

            item = Item(linha=n, empresa=emp, comp=comp_iso(comp), tipo_codigo=tipo.codigo, tipo_nome=tipo.nome)
            for grupo, rot in (("autor", "autorizadas"), ("canc", "canceladas")):
                q, v = valores.get(f"qtde_{grupo}"), valores.get(f"valor_{grupo}")
                if q is None and v is not None:
                    previa.avisar(f"Linha {n}: valor R$ {fmt_moeda(v)} sem quantidade ({rot}; {rotulo}) — importado com "
                                  "quantidade 0; corrija no lançamento", n, destaque=True)
                    q = 0
                elif v is None and q is not None:
                    previa.avisar(f"Linha {n}: quantidade {fmt_inteiro(q)} sem valor ({rot}; {rotulo}) — importado com "
                                  "valor 0,00; corrija no lançamento", n, destaque=True)
                    v = Decimal("0.00")
                setattr(item, f"qtde_{grupo}", q)
                setattr(item, f"valor_{grupo}", str(v) if v is not None else None)
                aq, av = valores.get(f"aj_qtde_{grupo}"), valores.get(f"aj_valor_{grupo}")
                setattr(item, f"aj_qtde_{grupo}", aq or 0)
                setattr(item, f"aj_valor_{grupo}", str(av if av is not None else Decimal("0.00")))
            previa.itens.append(item)

        for txt, linhas in tipos_desconhecidos.items():
            previa.avisar(f"TIPO desconhecido '{txt}' em {len(linhas)} linha(s) com números ({_listar_linhas(linhas)}) — "
                          "ignoradas", linhas[0], destaque=True)
        for c, linhas in futuras.items():
            previa.avisar(f"Competência futura {c} em {len(linhas)} linha(s) ({_listar_linhas(linhas)}) — ignoradas",
                          linhas[0])
    finally:
        wb_val.close()
        wb_form.close()
    previa.avisos.sort(key=lambda a: (a.linha or 0))
    return previa


def _listar_linhas(linhas: list[int], limite: int = 8) -> str:
    txt = ", ".join(str(x) for x in linhas[:limite])
    return f"linhas {txt}{'…' if len(linhas) > limite else ''}"


# ====================================================================== guarda entre requisições

def _pasta() -> Path:
    p = Path(current_app.config["DADOS_DIR"]) / "importacoes"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _limpar_antigos(pasta: Path, horas: int = 24) -> None:
    limite = time.time() - horas * 3600
    for f in pasta.iterdir():
        try:
            if f.is_file() and f.stat().st_mtime < limite:
                f.unlink()
        except OSError:  # pragma: no cover
            pass


def _caminho(token: str, ext: str) -> Path:
    if not token or not _TOKEN_OK.fullmatch(token):
        raise ValorInvalido("Prévia inválida ou expirada. Envie o arquivo de novo.")
    return _pasta() / f"{token}.{ext}"


def guardar_previa(previa: Previa) -> str:
    pasta = _pasta()
    _limpar_antigos(pasta)
    previa.token = previa.token or secrets.token_urlsafe(18)
    _caminho(previa.token, "json").write_text(json.dumps(previa.como_dict(), ensure_ascii=False), encoding="utf-8")
    return previa.token


def carregar_previa(token: str) -> Previa:
    p = _caminho(token, "json")
    if not p.exists():
        raise ValorInvalido("Prévia não encontrada ou expirada. Envie o arquivo de novo.")
    return Previa.de_dict(json.loads(p.read_text(encoding="utf-8")))


def descartar(token: str) -> None:
    for ext in ("json", "bin", "meta.json"):
        try:
            _caminho(token, ext).unlink(missing_ok=True)
        except ValorInvalido:
            return


def guardar_arquivo(conteudo: bytes, nome_arquivo: str) -> str:
    """Guarda um upload bruto (planilha de responsáveis) para confirmar sem reenviar."""
    pasta = _pasta()
    _limpar_antigos(pasta)
    token = secrets.token_urlsafe(18)
    _caminho(token, "bin").write_bytes(conteudo)
    _caminho(token, "meta.json").write_text(json.dumps({"arquivo": nome_arquivo}, ensure_ascii=False), encoding="utf-8")
    return token


def carregar_arquivo(token: str) -> tuple[bytes, str]:
    b, m = _caminho(token, "bin"), _caminho(token, "meta.json")
    if not b.exists():
        raise ValorInvalido("Arquivo não encontrado ou expirado. Envie a planilha de novo.")
    nome = json.loads(m.read_text(encoding="utf-8")).get("arquivo", "planilha.csv") if m.exists() else "planilha.csv"
    return b.read_bytes(), nome


# ====================================================================== gravação

def _valores_lanc(l: Lancamento) -> dict:
    return {c: getattr(l, c) for c in CAMPOS}


def _json_valores(d: dict) -> dict:
    return {k: (str(v) if isinstance(v, Decimal) else v) for k, v in d.items()}


def confirmar(previa: Previa, responsavel: Usuario, usuario: Usuario, nome_arquivo: str | None = None) -> Importacao:
    """Grava a prévia (uma transação; quem chama faz commit/rollback).

    - empresas novas: criadas com ``responsavel`` desde a primeira competência do arquivo e com os
      tipos que tiveram números;
    - empresas existentes: ganham os tipos que faltavam;
    - lançamento 'Bateu' nunca é sobrescrito; os demais são criados/atualizados (com auditoria);
    - colunas AJUSTE_* viram ajustes com motivo "Importado da planilha" (sem duplicar ao reimportar).
    """
    if responsavel is None or not responsavel.ativo:
        raise RegraViolada("Escolha o responsável pelas empresas novas.")
    if not previa.itens:
        raise RegraViolada("A planilha não tem nenhuma linha com números para importar.")
    nome_arquivo = nome_arquivo or previa.arquivo
    imp = Importacao(arquivo=nome_arquivo[:255], usuario_id=usuario.id if usuario else None,
                     responsavel_id=responsavel.id)
    db.session.add(imp)
    db.session.flush()

    tipos = {t.codigo: t for t in db.session.scalars(select(TipoDocumento))}
    existentes = _empresas_por_nome()
    inicio = min(i.competencia for i in previa.itens)
    avisos: list[dict] = [asdict(a) for a in previa.avisos]
    r = {
        "empresas_criadas": [], "empresas_criadas_ids": [], "empresas_existentes": 0,
        "tipos_adicionados": [], "lancamentos_criados": 0, "lancamentos_alterados": 0, "sem_mudanca": 0,
        "ignorados_bateu": 0, "ajustes_criados": 0, "ajustes_removidos": 0,
        "ajustes_ids": [], "alterados": [],
    }

    por_empresa: dict[str, list[Item]] = defaultdict(list)
    for i in previa.itens:
        por_empresa[normalizar_nome(i.empresa)].append(i)

    # ---- empresas
    empresas: dict[str, Empresa] = {}
    for chave, itens in por_empresa.items():
        tipo_ids = sorted({tipos[i.tipo_codigo].id for i in itens if i.tipo_codigo in tipos})
        emp = existentes.get(chave)
        if emp is None:
            emp = criar_empresa(nome_exibicao(itens[0].empresa), responsavel, usuario, tipo_ids=tipo_ids,
                                inicio=inicio, origem=ORIGEM_PLANILHA)
            r["empresas_criadas"].append(emp.nome)
            r["empresas_criadas_ids"].append(emp.id)
        else:
            r["empresas_existentes"] += 1
            atuais = {t.id for t in emp.tipos}
            if not set(tipo_ids) <= atuais:
                novos = [tipos[c].nome for c in tipos if tipos[c].id in set(tipo_ids) - atuais]
                atualizar_empresa(emp, usuario, tipo_ids=sorted(atuais | set(tipo_ids)))
                r["tipos_adicionados"].append(f"{emp.nome}: {', '.join(novos)}")
        empresas[chave] = emp

    # ---- lançamentos
    atuais = _lancamentos_existentes({e.id for e in empresas.values()})
    meses_tocados: set[tuple[int, date]] = set()
    for i in previa.itens:
        emp = empresas[normalizar_nome(i.empresa)]
        tipo = tipos[i.tipo_codigo]
        comp = i.competencia
        base = i.base()
        rotulo = f"{emp.nome} · {comp_curta(comp)} · {tipo.nome}"
        lanc = atuais.get((emp.id, comp, tipo.id))
        if lanc is not None and lanc.bateu:
            precisa = not _mesmos_numeros(lanc, base) or not _ajustes_iguais(lanc, i)
            if precisa:
                r["ignorados_bateu"] += 1
                avisos.append(asdict(Aviso(f"Linha {i.linha}: {rotulo} já está marcado como 'Bateu' — não alterado",
                                           i.linha, destaque=True)))
            else:
                r["sem_mudanca"] += 1
            continue

        if lanc is None:
            resp = responsavel_em(emp, comp)
            lanc = Lancamento(empresa_id=emp.id, competencia=comp, tipo_id=tipo.id, criado_por_id=usuario.id,
                              responsavel_id=resp.id if resp else None, importacao_id=imp.id, **base)
            db.session.add(lanc)
            db.session.flush()
            atuais[(emp.id, comp, tipo.id)] = lanc
            r["lancamentos_criados"] += 1
            auditoria.registrar("fat_lancamento", "criar", usuario, registro_id=lanc.id, empresa_id=emp.id,
                                depois=base, descricao=f"{rotulo} (importação #{imp.id}, linha {i.linha})")
            meses_tocados.add((emp.id, comp))
        elif not _mesmos_numeros(lanc, base):
            antes = _valores_lanc(lanc)
            for c in CAMPOS:
                setattr(lanc, c, base[c])
            lanc.alterado_por_id = usuario.id
            lanc.alterado_em = agora()
            r["lancamentos_alterados"] += 1
            r["alterados"].append({"id": lanc.id, "antes": _json_valores(antes)})
            auditoria.registrar("fat_lancamento", "alterar", usuario, registro_id=lanc.id, empresa_id=emp.id,
                                antes=antes, depois=base,
                                descricao=f"{rotulo}: números substituídos pela planilha (importação #{imp.id}, linha {i.linha})")
            meses_tocados.add((emp.id, comp))
        else:
            r["sem_mudanca"] += 1

        # ---- ajustes da planilha
        for grupo in (GRUPO_AUTOR, GRUPO_CANC):
            q, v = i.ajuste(grupo)
            importados = [a for a in lanc.ajustes if a.grupo == grupo and a.motivo == MOTIVO_IMPORTACAO]
            if len(importados) == 1 and importados[0].qtde == q and Decimal(importados[0].valor) == v:
                continue
            if not importados and q == 0 and v == 0:
                continue
            for a in importados:
                auditoria.registrar("fat_ajuste", "remover", usuario, registro_id=a.id, empresa_id=emp.id,
                                    antes={"grupo": a.grupo, "qtde": a.qtde, "valor": a.valor, "motivo": a.motivo},
                                    descricao=f"Ajuste importado substituído em {rotulo} (importação #{imp.id})")
                lanc.ajustes.remove(a)
                r["ajustes_removidos"] += 1
            if q == 0 and v == 0:
                continue
            aj = Ajuste(lancamento=lanc, grupo=grupo, qtde=q, valor=v, motivo=MOTIVO_IMPORTACAO,
                        observacao=f"Importação #{imp.id} — {nome_arquivo}, linha {i.linha}", criado_por_id=usuario.id)
            db.session.add(aj)
            db.session.flush()
            r["ajustes_criados"] += 1
            r["ajustes_ids"].append(aj.id)
            lanc.alterado_por_id = usuario.id
            lanc.alterado_em = agora()
            meses_tocados.add((emp.id, comp))
            auditoria.registrar("fat_ajuste", "criar", usuario, registro_id=aj.id, empresa_id=emp.id,
                                depois={"grupo": grupo, "qtde": q, "valor": v, "motivo": MOTIVO_IMPORTACAO},
                                descricao=f"Ajuste {'+' if q >= 0 else ''}{q} / R$ {fmt_moeda(v)} em {rotulo} "
                                          f"(importação #{imp.id})")

    emp_por_id = {e.id: e for e in empresas.values()}
    for eid, comp in meses_tocados:
        limpar_sem_movimento(emp_por_id[eid], comp, usuario)

    comps = previa.competencias
    r["avisos"] = avisos
    r["n_avisos"] = len(avisos)
    r["n_avisos_destaque"] = sum(1 for a in avisos if a.get("destaque"))
    r["competencias"] = [comp_iso(comps[0]), comp_iso(comps[-1])] if comps else []
    r["linhas_planilha"] = len(previa.itens)
    r["responsavel"] = responsavel.nome
    imp.resumo = json.dumps(r, ensure_ascii=False)
    auditoria.registrar("fat_importacao", "importar", usuario, registro_id=imp.id,
                        depois={k: v for k, v in r.items() if k not in ("avisos", "alterados", "ajustes_ids")},
                        descricao=f"Importação #{imp.id} de {nome_arquivo}: {r['lancamentos_criados']} lançamento(s) "
                                  f"novo(s), {r['lancamentos_alterados']} alterado(s), {r['ajustes_criados']} ajuste(s), "
                                  f"{len(r['empresas_criadas'])} empresa(s) criada(s)")
    db.session.flush()
    return imp


def _ajustes_iguais(lanc: Lancamento, item: Item) -> bool:
    for grupo in (GRUPO_AUTOR, GRUPO_CANC):
        q, v = item.ajuste(grupo)
        imp = [a for a in lanc.ajustes if a.grupo == grupo and a.motivo == MOTIVO_IMPORTACAO]
        if not imp and q == 0 and v == 0:
            continue
        if len(imp) != 1 or imp[0].qtde != q or Decimal(imp[0].valor) != v:
            return False
    return True


def resumo(imp: Importacao) -> dict:
    try:
        return json.loads(imp.resumo or "{}")
    except ValueError:
        return {}


def listar() -> list[Importacao]:
    return list(db.session.scalars(select(Importacao).order_by(Importacao.data.desc(), Importacao.id.desc())))


def desfazer(imp: Importacao, usuario: Usuario) -> dict:
    """Remove o que a importação criou (lançamentos não conferidos, ajustes, empresas que ficaram vazias)
    e devolve os números anteriores dos lançamentos que ela sobrescreveu."""
    if imp.desfeita:
        raise RegraViolada("Esta importação já foi desfeita.")
    r = resumo(imp)
    out = {"removidos": 0, "mantidos_bateu": 0, "restaurados": 0, "ajustes_removidos": 0, "empresas_removidas": 0}

    for l in list(db.session.scalars(select(Lancamento).where(Lancamento.importacao_id == imp.id))):
        rot = f"{l.empresa.nome} · {comp_curta(l.competencia)} · {l.tipo.nome}"
        if l.bateu:
            out["mantidos_bateu"] += 1
            continue
        auditoria.registrar("fat_lancamento", "remover", usuario, registro_id=l.id, empresa_id=l.empresa_id,
                            antes=_valores_lanc(l), descricao=f"{rot}: removido ao desfazer a importação #{imp.id}")
        db.session.delete(l)
        out["removidos"] += 1
    db.session.flush()

    for aid in r.get("ajustes_ids", []):
        aj = db.session.get(Ajuste, aid)
        if aj is None or aj.lancamento.bateu:
            continue
        lanc = aj.lancamento
        auditoria.registrar("fat_ajuste", "remover", usuario, registro_id=aj.id, empresa_id=lanc.empresa_id,
                            antes={"grupo": aj.grupo, "qtde": aj.qtde, "valor": aj.valor, "motivo": aj.motivo},
                            descricao=f"Ajuste removido ao desfazer a importação #{imp.id}")
        lanc.ajustes.remove(aj)
        out["ajustes_removidos"] += 1

    for alt in r.get("alterados", []):
        l = db.session.get(Lancamento, alt["id"])
        if l is None or l.bateu:
            continue
        antes = {c: (Decimal(v) if v is not None and c.startswith("valor") else v) for c, v in alt["antes"].items()}
        atual = _valores_lanc(l)
        for c in CAMPOS:
            setattr(l, c, antes.get(c))
        l.alterado_por_id = usuario.id if usuario else None
        l.alterado_em = agora()
        auditoria.registrar("fat_lancamento", "alterar", usuario, registro_id=l.id, empresa_id=l.empresa_id,
                            antes=atual, depois=antes,
                            descricao=f"{l.empresa.nome} · {comp_curta(l.competencia)} · {l.tipo.nome}: números "
                                      f"anteriores restaurados ao desfazer a importação #{imp.id}")
        out["restaurados"] += 1
    db.session.flush()

    for eid in r.get("empresas_criadas_ids", []):
        emp = db.session.get(Empresa, eid)
        if emp is None:
            continue
        resta = db.session.scalar(select(Lancamento.id).where(Lancamento.empresa_id == eid).limit(1))
        if resta is not None:
            continue
        for f in db.session.scalars(select(Fechamento).where(Fechamento.empresa_id == eid)):
            db.session.delete(f)
        auditoria.registrar("fat_empresa", "remover", usuario, registro_id=eid, empresa_id=eid,
                            antes={"nome": emp.nome}, descricao=f"Empresa {emp.nome} removida ao desfazer a importação #{imp.id}")
        db.session.delete(emp)
        out["empresas_removidas"] += 1

    imp.desfeita = True
    auditoria.registrar("fat_importacao", "desfazer", usuario, registro_id=imp.id, depois=out,
                        descricao=f"Importação #{imp.id} ({imp.arquivo}) desfeita: {out['removidos']} lançamento(s) "
                                  f"removido(s), {out['restaurados']} restaurado(s), {out['mantidos_bateu']} mantido(s) por "
                                  "já estarem conferidos")
    db.session.flush()
    return out

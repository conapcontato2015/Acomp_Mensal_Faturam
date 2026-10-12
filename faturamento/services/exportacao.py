"""Exportações da conferência (Excel).

Regra: entradas e saídas nunca se somam — os subtotais são por empresa + sentido.
Valores exportados = totais com ajustes (as colunas de ajuste mostram só a parte ajustada).
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import ENTRADA
from .resumo import CartaoEmpresa

MARINHO = "202F5A"
CINZA = "F4F6FA"
FMT_INT = "#,##0"
FMT_MOEDA = "#,##0.00"

COLUNAS = [
    # título, largura, formato
    ("Empresa", 42, None),
    ("Responsável", 16, None),
    ("Sentido", 10, None),
    ("Tipo", 18, None),
    ("Qtde autorizada", 15, FMT_INT),
    ("Valor autorizado", 18, FMT_MOEDA),
    ("Qtde cancelada", 15, FMT_INT),
    ("Valor cancelado", 16, FMT_MOEDA),
    ("Ajuste qtde", 12, FMT_INT),
    ("Ajuste valor", 15, FMT_MOEDA),
    ("Bateu", 8, None),
]


def nome_arquivo_conferencia(comp: date) -> str:
    return f"conferencia-{comp.year:04d}-{comp.month:02d}.xlsx"


def _sentido_nome(s: str) -> str:
    return "Entrada" if s == ENTRADA else "Saída"


def planilha_conferencia(comp: date, cartoes: list[CartaoEmpresa]) -> bytes:
    """Monta o .xlsx da conferência e devolve os bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Conferência {comp.month:02d}-{comp.year:04d}"

    cab_font = Font(bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor=MARINHO)
    sub_font = Font(bold=True)
    sub_fill = PatternFill("solid", fgColor=CINZA)
    linha_fina = Border(top=Side(style="thin", color="DCE0EA"))

    ws.append([c[0] for c in COLUNAS])
    for i, (_, largura, _fmt) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=i)
        cel.font = cab_font
        cel.fill = cab_fill
        cel.alignment = Alignment(horizontal="center" if i > 4 else "left", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    def formatar(row: int) -> None:
        for i, (_, _l, fmt) in enumerate(COLUNAS, start=1):
            if fmt:
                ws.cell(row=row, column=i).number_format = fmt

    for c in cartoes:
        resp = c.responsavel_nome or ""
        for b in c.blocos:
            sentido = _sentido_nome(b.sentido)
            for l in b.linhas:
                ws.append([
                    c.empresa.nome, resp, sentido, l.tipo.nome,
                    l.total_qtde_autor, l.total_valor_autor, l.total_qtde_canc, l.total_valor_canc,
                    l.ajuste_qtde_autor, l.ajuste_valor_autor, "Sim" if l.bateu else "Não",
                ])
                formatar(ws.max_row)
            ws.append([
                c.empresa.nome, resp, sentido, f"Subtotal {'entradas' if b.sentido == ENTRADA else 'saídas'}",
                b.qtde_autor, b.valor_autor, b.qtde_canc, b.valor_canc,
                sum(l.ajuste_qtde_autor for l in b.linhas), sum((l.ajuste_valor_autor for l in b.linhas), 0) or 0,
                f"{sum(1 for l in b.linhas if l.bateu)} de {len(b.linhas)}",
            ])
            r = ws.max_row
            formatar(r)
            for i in range(1, len(COLUNAS) + 1):
                cel = ws.cell(row=r, column=i)
                cel.font = sub_font
                cel.fill = sub_fill
                cel.border = linha_fina

    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ws.max_row}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
